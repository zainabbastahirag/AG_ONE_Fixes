#!/usr/bin/env python3
"""Generate a structured Word document for the Nexa tech team alignment session."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).parent / "Nexa-Tech-Team-Session-Summary.docx"

# Brand colours
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x00, 0x7A, 0x87)
LIGHT_GREY = "F2F4F7"
WHITE = "FFFFFF"
HEADER_BG = "1F3A5F"
ACCENT_BG = "E8F4F6"


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def style_paragraph(paragraph, text: str, *, bold=False, size=11, color=None, space_after=6):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    return run


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = NAVY if level == 1 else TEAL


def add_body(doc: Document, text: str, bold=False) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, text, bold=bold)


def format_header_row(table, labels: list[str]) -> None:
    row = table.rows[0]
    for idx, label in enumerate(labels):
        cell = row.cells[idx]
        set_cell_shading(cell, HEADER_BG)
        set_cell_margins(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = style_paragraph(p, label, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), space_after=0)


def fill_data_row(table, row_idx: int, values: list[str], bold_first=False, alt=False) -> None:
    row = table.rows[row_idx]
    bg = LIGHT_GREY if alt else WHITE
    for col_idx, value in enumerate(values):
        cell = row.cells[col_idx]
        set_cell_shading(cell, bg)
        set_cell_margins(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        style_paragraph(
            p,
            value,
            bold=(bold_first and col_idx == 0),
            size=10,
            space_after=0,
        )


def add_styled_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    format_header_row(table, headers)
    for i, row_data in enumerate(rows, start=1):
        fill_data_row(table, i, row_data, bold_first=True, alt=(i % 2 == 0))
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def build_document() -> Document:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, "Nexa", bold=True, size=28, color=NAVY, space_after=4)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(
        subtitle,
        "Tech Team Alignment Session",
        bold=True,
        size=16,
        color=TEAL,
        space_after=12,
    )

    add_styled_table(
        doc,
        ["Field", "Details"],
        [
            ["Purpose", "Align on Nexa's current state, priorities, and Product Team inputs needed to support continuous improvement within the AG ONE ecosystem."],
            ["Reference Material", "UAT - 28th June 2026.xlsx (Nexa tab)"],
            ["Session Duration", "60–90 minutes"],
            ["Document Version", "1.0 | July 2026"],
        ],
        col_widths=[1.6, 4.9],
    )

    # Discussion points
    add_heading(doc, "Key Discussion Points", level=1)
    add_styled_table(
        doc,
        ["#", "Topic", "Discussion Focus"],
        [
            ["1", "Current capabilities & product vision", "Review what Nexa can do today; assess alignment with product vision and AG ONE strategy."],
            ["2", "UAT review (28 June 2026)", "Walk through Nexa tab items; confirm status, severity, and ownership for open items."],
            ["3", "Hybrid knowledge model", "How Nexa combines KB-driven knowledge with dynamic data sources; gaps, risks, and opportunities."],
            ["4", "Lead conversion & user journeys", "Navigation guidance, lead generation, sign-up, product discovery, and purchase handoffs."],
            ["5", "Response formats, behavior & UX", "Tone, structure, interaction patterns, and overall end-user experience quality."],
            ["6", "Troubleshooting, support & guidance", "In-conversation support, ticket submission, escalation paths, and self-serve boundaries."],
        ],
        col_widths=[0.35, 1.75, 4.4],
    )

    # Product team inputs
    add_heading(doc, "Information Required from the Product Team", level=1)
    add_styled_table(
        doc,
        ["Item", "What We Need", "Owner"],
        [
            ["Use cases", "Use-case template and populated examples", "Product Team"],
            ["Additional requirements", "Any new or changed requirements not yet captured", "Product Team"],
            ["Considerations", "Constraints, compliance, branding, or market-specific needs", "Product Team"],
            ["Relevant documentation", "PRDs, journey maps, KB ownership, or other reference material", "Product Team"],
        ],
        col_widths=[1.5, 3.5, 1.5],
    )

    # Expected outcomes
    add_heading(doc, "Expected Outcomes", level=1)
    add_styled_table(
        doc,
        ["Outcome", "Description"],
        [
            ["Alignment", "Shared understanding of Nexa's role and objectives within the AG ONE ecosystem."],
            ["Prioritization", "Identified improvement opportunities and a ranked enhancement backlog."],
            ["Clarity on inputs", "Clear list of Product Team deliverables required for continuous improvement and scalability."],
        ],
        col_widths=[1.5, 5.0],
    )

    # Attendees
    add_heading(doc, "Attendees", level=1)
    add_body(doc, "Please extend this invitation to any relevant Nexa team members as needed.")
    add_styled_table(
        doc,
        ["Role / Team", "Why They Should Attend"],
        [
            ["Nexa engineering / platform", "Technical capabilities, architecture, and implementation constraints."],
            ["Product (AG ONE / Nexa)", "Vision alignment, use cases, and requirement prioritization."],
            ["QA / UAT owners", "Review of 28 June UAT findings and validation status."],
            ["Support / operations", "Ticket flows, escalation paths, and operational guidance needs."],
        ],
        col_widths=[2.0, 4.5],
    )

    # Agenda
    add_heading(doc, "Suggested Agenda", level=1)
    add_styled_table(
        doc,
        ["Time", "Topic", "Objective"],
        [
            ["5 min", "Objectives & expected outcomes", "Set context and agree on session goals."],
            ["15 min", "Current capabilities vs. product vision", "Baseline understanding of where Nexa stands today."],
            ["20 min", "UAT Nexa tab — review & triage", "Review open items; assign severity and ownership."],
            ["15 min", "Hybrid knowledge model", "Discuss KB + dynamic data approach and gaps."],
            ["15 min", "Lead conversion, UX & support flows", "Align on user journeys and experience standards."],
            ["10 min", "Product inputs & next steps", "Confirm deliverables, owners, and follow-up cadence."],
        ],
        col_widths=[0.75, 2.25, 3.5],
    )

    # Pre-read
    add_heading(doc, "Pre-read / Bring to the Session", level=1)
    add_styled_table(
        doc,
        ["Check", "Material", "Status"],
        [
            ["☐", "UAT - 28th June 2026.xlsx (Nexa tab)", "Required"],
            ["☐", "Current Nexa capability overview (if available)", "Recommended"],
            ["☐", "Use-case template (Product Team)", "Required"],
            ["☐", "List of open questions or blockers from engineering", "Recommended"],
        ],
        col_widths=[0.5, 4.0, 1.0],
    )

    # Next steps
    add_heading(doc, "Next Steps (to Confirm in Session)", level=1)
    add_styled_table(
        doc,
        ["Action Item", "Owner", "Target Date"],
        [
            ["Agreed priority list from UAT and capability gaps", "TBD", "TBD"],
            ["Owners and target dates for top items", "TBD", "TBD"],
            ["Product deliverables and due dates (use cases, docs)", "Product Team", "TBD"],
            ["Follow-up cadence (e.g. weekly sync or milestone review)", "TBD", "TBD"],
        ],
        col_widths=[3.2, 1.4, 1.4],
    )

    # Footer note
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(
        note,
        "AG ONE  |  Nexa Alignment Session  |  Confidential — Internal Use",
        size=9,
        color=RGBColor(0x66, 0x66, 0x66),
        space_after=0,
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()

"""Generate diagrams + a Word (.docx) implementation guide for adding
Global + Tenant scope to Master Data."""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = "/opt/cursor/artifacts"
IMG = os.path.join(OUT, "_img")
os.makedirs(IMG, exist_ok=True)

# ---- palette ----
BLUE = "#2563eb"
GREEN = "#16a34a"
AMBER = "#d97706"
GREY = "#475569"
LIGHT = "#e2e8f0"
GLOBAL_C = "#0ea5e9"
TENANT_C = "#22c55e"


def box(ax, x, y, w, h, text, fc, ec="#1e293b", tc="white", fs=11, bold=True):
    p = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                       linewidth=1.4, edgecolor=ec, facecolor=fc, mutation_aspect=1)
    ax.add_patch(p)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            fontsize=fs, color=tc, weight="bold" if bold else "normal", wrap=True)


def arrow(ax, x1, y1, x2, y2, color=GREY, style="-|>", lw=1.6, ls="-"):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle=style,
                 mutation_scale=16, color=color, lw=lw, linestyle=ls,
                 shrinkA=2, shrinkB=2))


def container(ax, x, y, w, h, title, fc, tc="white", fs=11):
    p = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.08",
                       linewidth=1.4, edgecolor="#1e293b", facecolor=fc)
    ax.add_patch(p)
    ax.text(x + w / 2, y + h - 0.22, title, ha="center", va="center",
            fontsize=fs, color=tc, weight="bold")


def new_ax(w=10, h=5.5):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, h)
    ax.axis("off")
    return fig, ax


# =====================================================================
# Diagram 1: Current vs Target scope model
# =====================================================================
fig, ax = new_ax(11, 5.2)
ax.set_ylim(0, 5.2)
ax.text(2.6, 5.0, "BEFORE  (Tenant scope only)", ha="center", fontsize=13, weight="bold", color=GREY)
ax.text(8.0, 5.0, "AFTER  (Global + Tenant)", ha="center", fontsize=13, weight="bold", color=BLUE)
# before
box(ax, 0.4, 3.4, 1.8, 0.9, "Tenant A", TENANT_C, fs=10)
box(ax, 0.4, 2.0, 1.8, 0.9, "Tenant B", TENANT_C, fs=10)
box(ax, 3.0, 3.4, 2.0, 0.9, "Master Data\nTenantId = A", GREY, fs=9)
box(ax, 3.0, 2.0, 2.0, 0.9, "Master Data\nTenantId = B", GREY, fs=9)
arrow(ax, 2.2, 3.85, 3.0, 3.85)
arrow(ax, 2.2, 2.45, 3.0, 2.45)
ax.plot([5.4, 5.4], [1.4, 4.6], color=LIGHT, lw=2)
# after
box(ax, 6.0, 4.05, 3.4, 0.85, "GLOBAL Master Data  (TenantId = NULL)", GLOBAL_C, fs=9)
box(ax, 5.8, 2.55, 1.7, 0.85, "Tenant A", TENANT_C, fs=9)
box(ax, 8.0, 2.55, 1.7, 0.85, "Tenant B", TENANT_C, fs=9)
box(ax, 5.8, 1.2, 1.7, 0.85, "MD  Tenant=A", GREY, fs=8)
box(ax, 8.0, 1.2, 1.7, 0.85, "MD  Tenant=B", GREY, fs=8)
arrow(ax, 6.65, 2.55, 6.65, 2.05)
arrow(ax, 8.85, 2.55, 8.85, 2.05)
arrow(ax, 7.0, 4.05, 6.65, 3.4, color=GLOBAL_C, ls="--")
arrow(ax, 8.4, 4.05, 8.85, 3.4, color=GLOBAL_C, ls="--")
ax.text(7.7, 0.75, "Each tenant sees:  Global  +  its own Tenant data",
        ha="center", fontsize=9, style="italic", color=BLUE)
fig.tight_layout()
p1 = os.path.join(IMG, "d1_scope_model.png")
fig.savefig(p1, dpi=150, bbox_inches="tight")
plt.close(fig)

# =====================================================================
# Diagram 2: Two-schema architecture (same DB)
# =====================================================================
fig, ax = new_ax(10, 5.4)
ax.set_ylim(0, 5.4)
# outer DB
outer = FancyBboxPatch((0.4, 0.4), 9.2, 4.4, boxstyle="round,pad=0.02,rounding_size=0.1",
                       linewidth=2, edgecolor=BLUE, facecolor="#f1f5f9")
ax.add_patch(outer)
ax.text(5.0, 4.5, "ONE PHYSICAL DATABASE", ha="center", fontsize=12, weight="bold", color=BLUE)
# ag schema
container(ax, 0.9, 2.2, 3.8, 1.9, "schema:  ag   (metadata)", GREY, fs=11)
box(ax, 1.15, 3.1, 3.3, 0.5, "MasterDataDefinition", "#64748b", fs=9)
box(ax, 1.15, 2.45, 3.3, 0.5, "MasterDataField", "#94a3b8", tc="#0f172a", fs=9)
ax.text(2.8, 1.95, "unchanged", ha="center", fontsize=9, style="italic", color=GREEN)
# products schema
container(ax, 5.3, 2.2, 3.8, 1.9, "schema:  products   (data)", GREY, fs=11)
box(ax, 5.55, 3.1, 3.3, 0.5, "MasterDataRecord", GREEN, fs=9)
box(ax, 5.55, 2.45, 3.3, 0.5, "+ Scope   + TenantId(NULL)", AMBER, fs=8)
ax.text(7.2, 1.95, "only table that changes", ha="center", fontsize=9, style="italic", color=AMBER)
# relation
arrow(ax, 4.7, 3.35, 5.55, 3.35, color=BLUE)
ax.text(5.12, 3.6, "DefinitionId", ha="center", fontsize=8, color=BLUE)
ax.text(5.0, 1.3, "Same DB, different schemas — layout preserved",
        ha="center", fontsize=10, style="italic", color=GREY)
fig.tight_layout()
p2 = os.path.join(IMG, "d2_architecture.png")
fig.savefig(p2, dpi=150, bbox_inches="tight")
plt.close(fig)

# =====================================================================
# Diagram 3: Read / merge resolution flow
# =====================================================================
fig, ax = new_ax(10, 5.6)
ax.set_ylim(0, 5.6)
box(ax, 3.7, 4.8, 2.6, 0.7, "Tenant request\n(DefinitionId)", BLUE, fs=10)
box(ax, 3.4, 3.5, 3.2, 0.8, "WHERE Scope=Global\nOR TenantId = current", GREY, fs=9)
arrow(ax, 5.0, 4.8, 5.0, 4.3)
box(ax, 0.8, 2.1, 3.0, 0.8, "Global rows\n(TenantId NULL)", GLOBAL_C, fs=9)
box(ax, 6.2, 2.1, 3.0, 0.8, "Tenant rows\n(TenantId = current)", TENANT_C, fs=9)
arrow(ax, 4.2, 3.5, 2.3, 2.9)
arrow(ax, 5.8, 3.5, 7.7, 2.9)
box(ax, 3.2, 0.7, 3.6, 0.85, "MERGE by natural key\nTenant overrides Global", AMBER, fs=9)
arrow(ax, 2.3, 2.1, 4.2, 1.55, color=GLOBAL_C)
arrow(ax, 7.7, 2.1, 5.8, 1.55, color=TENANT_C)
ax.text(5.0, 0.25, "Result = one merged list the tenant sees",
        ha="center", fontsize=10, style="italic", color=BLUE)
fig.tight_layout()
p3 = os.path.join(IMG, "d3_read_merge.png")
fig.savefig(p3, dpi=150, bbox_inches="tight")
plt.close(fig)

# =====================================================================
# Diagram 4: Implementation sequence
# =====================================================================
fig, ax = new_ax(10, 3.0)
ax.set_ylim(0, 3.0)
steps = ["1. Migration\nScope + nullable\nTenantId", "2. Update read\nquery filter",
         "3. Merge /\noverride rule", "4. Guarded\nGlobal writes", "5. Seed global\n+ rollout"]
cols = [GREY, BLUE, AMBER, GREEN, GLOBAL_C]
x = 0.3
for s, c in zip(steps, cols):
    box(ax, x, 1.0, 1.7, 1.1, s, c, fs=8)
    if x > 0.3:
        arrow(ax, x - 0.2, 1.55, x, 1.55)
    x += 1.95
fig.tight_layout()
p4 = os.path.join(IMG, "d4_sequence.png")
fig.savefig(p4, dpi=150, bbox_inches="tight")
plt.close(fig)

# =====================================================================
# Build the Word document
# =====================================================================
doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def para(text, italic=False, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def numbered(text):
    doc.add_paragraph(text, style="List Number")


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    # light shaded background via table would be heavier; keep monospace
    return p


def img(path, width=6.3):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)


# ----- Title page -----
t = doc.add_heading("Master Data: Adding Global + Tenant Scope", level=0)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Solution Design & Implementation Guide")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x25, 0x63, 0xeb)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Status: Draft for review    |    Date: 2026-07-14    |    Scope: minimal-impact change").italic = True

doc.add_paragraph()
note = doc.add_paragraph()
note.add_run("Note: ").bold = True
note.add_run("The described Master Data system is not in the current repository, so this "
             "guide is written from your description + multi-tenant best practice. Confirm the "
             "assumptions before implementing.").italic = True

# ----- 1. Goal -----
h("1. Goal (in one line)", 1)
para("Let master data be either GLOBAL (one copy shared by every tenant) or TENANT "
     "(owned by a single tenant), with the smallest possible change and full backward "
     "compatibility.")
bullet("Global data = maintained once, centrally (e.g. countries, currencies, standard categories).")
bullet("Tenant data = today's behaviour, unchanged.")
bullet("Each tenant sees: Global data + its own Tenant data, merged.")

# ----- 2. Before vs after -----
h("2. Concept: before vs. after", 1)
img(p1, 6.6)
caption("Figure 1 — Today only Tenant scope exists; after the change, Global data is shared "
        "and merged with each tenant's own data.")

# ----- 3. Architecture -----
h("3. Where it lives (architecture)", 1)
para("Metadata (definitions & fields) stays in the ag schema and does NOT change. Only the "
     "MasterDataRecord table in the products schema gains scope. Same database, same two-schema "
     "layout.")
img(p2, 6.6)
caption("Figure 2 — One physical database, two schemas. Only products.MasterDataRecord changes.")

# ----- 4. Core idea -----
h("4. Core idea (why it is minimal)", 1)
para("Add one Scope column and make TenantId nullable. Everything else derives from these two "
     "columns:")
code("Scope = Global  ->  TenantId IS NULL      (shared by all tenants)\n"
     "Scope = Tenant  ->  TenantId = <tenant>    (today's behaviour)")
para("The new column defaults to Tenant, so every existing row auto-migrates to Tenant scope — "
     "no data backfill, no behaviour change.")

# ----- 5. Read/merge -----
h("5. How a tenant reads data", 1)
img(p3, 5.6)
caption("Figure 3 — One query returns Global + Tenant rows; a merge step lets a tenant row "
        "override a global row with the same key.")

# ----- 6. Implementation guide -----
h("6. Implementation guide (step by step)", 1)
img(p4, 6.6)
caption("Figure 4 — Five implementation steps, backward compatible at every stage.")

h("Step 1 — Database migration (products.MasterDataRecord)", 2)
code("ALTER TABLE products.MasterDataRecord\n"
     "    ADD Scope TINYINT NOT NULL\n"
     "    CONSTRAINT DF_MDR_Scope DEFAULT (2);   -- 2 = Tenant (existing rows)\n\n"
     "ALTER TABLE products.MasterDataRecord\n"
     "    ALTER COLUMN TenantId UNIQUEIDENTIFIER NULL;\n\n"
     "ALTER TABLE products.MasterDataRecord\n"
     "    ADD CONSTRAINT CK_MDR_Scope_Tenant CHECK (\n"
     "        (Scope = 1 AND TenantId IS NULL) OR\n"
     "        (Scope = 2 AND TenantId IS NOT NULL));")
para("Uniqueness — natural key unique per scope bucket (filtered indexes):", italic=True)
code("CREATE UNIQUE INDEX UX_MDR_Tenant_Key\n"
     "  ON products.MasterDataRecord (DefinitionId, TenantId, Code) WHERE Scope = 2;\n\n"
     "CREATE UNIQUE INDEX UX_MDR_Global_Key\n"
     "  ON products.MasterDataRecord (DefinitionId, Code) WHERE Scope = 1;")

h("Step 2 — Model + read filter (EF Core example)", 2)
code("public enum MasterDataScope : byte { Global = 1, Tenant = 2 }\n\n"
     "public class MasterDataRecord {\n"
     "    public Guid Id { get; set; }\n"
     "    public int DefinitionId { get; set; }\n"
     "    public MasterDataScope Scope { get; set; } = MasterDataScope.Tenant;\n"
     "    public Guid? TenantId { get; set; }   // NULL when Global\n"
     "}")
para("Change the read in ONE place — the tenant query filter:", italic=True)
code("modelBuilder.Entity<MasterDataRecord>().HasQueryFilter(r =>\n"
     "    r.Scope == MasterDataScope.Global ||\n"
     "    r.TenantId == _tenantProvider.CurrentTenantId);")

h("Step 3 — Merge / override rule", 2)
para("If Global and Tenant both have the same natural key, the tenant row wins for that tenant:")
code("var merged = rows\n"
     "    .GroupBy(r => r.Code)\n"
     "    .Select(g => g.OrderBy(r => r.Scope).Last())  // Tenant(2) after Global(1)\n"
     "    .ToList();")

h("Step 4 — Guarded Global writes", 2)
bullet("Create Tenant record: unchanged (Scope defaults to Tenant, TenantId = current).")
bullet("Create Global record: Scope = Global, TenantId = NULL — allowed only for a "
       "platform/admin permission (e.g. MasterData.ManageGlobal). Enforce server-side.")
bullet("Read is unchanged for tenant users — they see the merged view with no new permission.")

h("Step 5 — Seed & rollout", 2)
numbered("Deploy the migration (default Scope = Tenant keeps current behaviour).")
numbered("Deploy code with the updated read filter + merge (optionally behind a feature flag).")
numbered("Enable Global authoring for admins and seed initial global data.")
para("Rollback is safe: old read path simply stops returning Global rows; tenant data is "
     "untouched.", italic=True)

# ----- 7. Impact table -----
h("7. Impact summary (how little changes)", 1)
rows = [
    ("Area", "Change", "Size"),
    ("products.MasterDataRecord", "+Scope col, nullable TenantId, 1 check + 2 indexes", "Small"),
    ("ag.* metadata schema", "None", "None"),
    ("Read path", "1 predicate in the central query filter", "Small"),
    ("Merge / override", "1 grouping step", "Small"),
    ("Write path", "Tenant unchanged; add guarded global create", "Small"),
    ("Security", "1 new permission + server guard", "Small"),
    ("Data migration", "Column default auto-backfills to Tenant", "Trivial"),
]
table = doc.add_table(rows=0, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(rows):
    cells = table.add_row().cells
    for j, val in enumerate(row):
        cells[j].text = val
        if i == 0:
            for pr in cells[j].paragraphs:
                for rn in pr.runs:
                    rn.bold = True

# ----- 8. Decisions & assumptions -----
h("8. Decisions to confirm", 1)
numbered("\"AG one\" = the ag metadata schema? (assumed yes)")
numbered("Collision policy: tenant-overrides-global (assumed) OR global is authoritative "
         "(reject tenant duplicates)?")
numbered("Scope only DATA (recommended, smallest) OR also DEFINITIONS (tenants define their own types)?")
numbered("Which role/permission may author Global data?")
numbered("Is there a single central tenant query filter to amend, or per-query filters?")

doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.add_run("End of document").italic = True

out_path = os.path.join(OUT, "MasterData-Global-Tenant-Scope-Implementation-Guide.docx")
doc.save(out_path)
print("SAVED", out_path)
print("SIZE", os.path.getsize(out_path), "bytes")
